"""ออเดอร์: รายการ, รายละเอียด, ตะกร้า, ชำระเงิน, ใบเสร็จ, Export"""

from ..forms import CustomerQRPaymentForm
from ._common import *  # noqa: F403
from ._common import _promptpay_qr_for, _sale_rows, _user_can_access_order


@login_required
def order_list(request):
    """แสดงรายการออเดอร์ทั้งหมด + เปิดออเดอร์ใหม่ + ค้นหา/กรองสถานะ/ประเภท"""
    query = request.GET.get('q', '').strip()
    status_filter = request.GET.get('status', '')
    type_filter = request.GET.get('type', '')
    table_from_qr = request.GET.get('table', '')

    if request.method == 'POST':
        # ผู้ใช้ทั่วไป: อนุญาตเปิดได้ครั้งละ 1 ออเดอร์ จนกว่าจะชำระ/ยกเลิกแล้วถึงจะเปิดใหม่ได้
        if request.user.is_authenticated and not request.user.is_staff:
            active_order = Order.objects.filter(
                created_by=request.user
            ).exclude(status__in=['paid', 'cancelled', 'refunded', 'completed']).first()
            if active_order:
                messages.warning(
                    request,
                    f'คุณมีออเดอร์ที่ยังไม่ชำระ ({active_order.display_code} {active_order.table_display}) '
                    f'กรุณาชำระเงินหรือจัดการออเดอร์ปัจจุบันก่อนจึงจะเปิดออเดอร์ใหม่ได้',
                )
                return redirect('order_detail', pk=active_order.pk)

        form = OrderForm(request.POST)
        if form.is_valid():
            order = form.save(commit=False)
            if request.user.is_authenticated:
                order.created_by = request.user
            if order.order_type == 'delivery':
                order.status = 'pending'
            order.save()
            messages.success(
                request,
                f'เปิดออเดอร์ {order.display_code} {order.table_display} เรียบร้อย'
                + (' — อย่าลืมกรอกข้อมูลจัดส่ง' if order.order_type == 'delivery' else ''),
            )
            return redirect('order_detail', pk=order.pk)
        messages.error(request, 'กรุณาตรวจสอบข้อมูลให้ถูกต้อง (ทานในร้านต้องกรอกเลขโต๊ะ 1-999)')
    else:
        form = OrderForm(
            initial={'table_number': table_from_qr, 'order_type': type_filter}
            if table_from_qr.isdigit() and type_filter == 'delivery'
            else ({'table_number': table_from_qr} if table_from_qr.isdigit() else {})
        )

    if request.user.is_authenticated and not request.user.is_staff:
        orders = Order.objects.select_related('delivery', 'created_by').filter(created_by=request.user)
    else:
        orders = Order.objects.select_related('delivery', 'created_by').all()

    if query:
        orders = orders.filter(
            Q(table_number__icontains=query)
            | Q(daily_code__icontains=query)
            | Q(delivery__recipient_name__icontains=query)
            | Q(delivery__phone__icontains=query)
        )
    if status_filter:
        orders = orders.filter(status=status_filter)
    if type_filter:
        orders = orders.filter(order_type=type_filter)

    paginator = Paginator(orders, 8)
    page_obj = paginator.get_page(request.GET.get('page'))

    status_counts = {
        'open': Order.objects.filter(status='open').count(),
        'preparing': Order.objects.filter(status='preparing').count(),
        'ready': Order.objects.filter(status='ready').count(),
        'served': Order.objects.filter(status='served').count(),
        'pending': Order.objects.filter(status='pending').count(),
        'confirmed': Order.objects.filter(status='confirmed').count(),
        'completed': Order.objects.filter(status='completed').count(),
        'refunded': Order.objects.filter(status='refunded').count(),
    }

    context = {
        'page_obj': page_obj,
        'form': form,
        'query': query,
        'status_filter': status_filter,
        'type_filter': type_filter,
        'status_counts': status_counts,
        'STATUS_CHOICES': Order.STATUS_CHOICES,
        'ORDER_TYPE_CHOICES': Order.ORDER_TYPE_CHOICES,
        'delivery_count': Order.objects.filter(order_type='delivery').count(),
        'dine_in_count': Order.objects.filter(order_type='dine_in').count(),
        'pending_delivery_count': Order.objects.filter(
            order_type='delivery', status__in=['pending', 'confirmed', 'preparing', 'ready', 'served']
        ).count(),
        'has_active_order': (
            request.user.is_authenticated
            and not request.user.is_staff
            and Order.objects.filter(created_by=request.user)
            .exclude(status__in=['paid', 'cancelled', 'refunded', 'completed']).exists()
        ),
        'active_order': (
            Order.objects.filter(created_by=request.user)
            .exclude(status__in=['paid', 'cancelled', 'refunded', 'completed']).first()
            if request.user.is_authenticated and not request.user.is_staff
            else None
        ),
    }
    return render(request, 'order/order_list.html', context)


@login_required
def order_detail(request, pk):
    """หน้าจัดการออเดอร์: เพิ่ม/ลบรายการ (แบ่งหน้า) สรุปราคารวม + ข้อมูลจัดส่ง (เดลิเวอร์รี่)"""
    order = get_object_or_404(
        Order.objects.prefetch_related('items', 'items__menu', 'items__addons').select_related('delivery'),
        pk=pk,
    )
    if not _user_can_access_order(request.user, order):
        messages.error(request, 'คุณไม่มีสิทธิ์เข้าถึงออเดอร์นี้')
        return redirect('order_list')
    form = OrderItemForm()

    try:
        delivery = order.delivery
    except Delivery.DoesNotExist:
        delivery = None

    # บันทึก/แก้ไขข้อมูลจัดส่ง (เฉพาะออเดอร์เดลิเวอร์รี่)
    if request.method == 'POST' and 'delivery_submit' in request.POST:
        delivery_form = DeliveryForm(request.POST, instance=delivery)
        if delivery_form.is_valid():
            delivery_obj = delivery_form.save(commit=False)
            delivery_obj.order = order
            # ระบบคำนวณระยะทาง + ค่าส่งอัตโนมัติภายใน Delivery.save() ถ้ามีพิกัดครบ
            delivery_obj.save()
            messages.success(request, 'บันทึกข้อมูลจัดส่งเรียบร้อยแล้ว')
            return redirect('order_detail', pk=order.pk)
        messages.error(request, 'กรุณาตรวจสอบข้อมูลจัดส่งให้ถูกต้อง')
    else:
        delivery_form = DeliveryForm(instance=delivery) if delivery else DeliveryForm()

    paginator = Paginator(order.items.all(), 5)  # แสดงรายการหน้าละ 5 แถว
    page_obj = paginator.get_page(request.GET.get('page'))

    store = StoreLocation.get_setting()
    context = {
        'order': order,
        'form': form,
        'delivery': delivery,
        'delivery_form': delivery_form,
        'delivery_status_choices': Delivery.DELIVERY_STATUS_CHOICES,
        'page_obj': page_obj,
        'store_lat': store.lat or 16.4539,
        'store_lng': store.lng or 99.5123,
        'available_menus': MenuItem.objects.filter(available=True).select_related('category_fk').order_by('category_fk', 'name'),
    }
    return render(request, 'order/order_detail.html', context)


@require_POST
@login_required
def order_add_item(request, pk):
    """เพิ่มเมนูลงออเดอร์
    - ถ้ามีรายการเดิม (เมนู + เครื่องเคียงชุดเดียวกัน) ให้รวมจำนวน
    - ถ้าเครื่องเคียงต่างกันให้แยกเป็นรายการใหม่
    - redirect กลับไปหน้าเดิมที่ user อยู่
    """
    order = get_object_or_404(Order, pk=pk)
    if not _user_can_access_order(request.user, order):
        messages.error(request, 'คุณไม่มีสิทธิ์แก้ไขออเดอร์นี้')
        return redirect('order_list')
    back_page = request.POST.get('page', '') or request.GET.get('page', '')

    if order.status in ('paid', 'cancelled', 'refunded', 'completed'):
        messages.error(request, 'ไม่สามารถเพิ่มรายการในออเดอร์ที่ชำระแล้ว ยกเลิกแล้ว คืนเงินแล้ว หรือเสร็จสิ้นแล้ว')
        return redirect('order_detail', pk=order.pk)

    form = OrderItemForm(request.POST)
    selected = ''
    if form.is_valid():
        menu = form.cleaned_data['menu']
        quantity = form.cleaned_data['quantity']
        notes = form.cleaned_data.get('notes', '')
        addons = list(form.cleaned_data.get('addons') or [])
        addons_pk = {a.pk for a in addons}

        # หารายการเดิมที่มีเมนูและเครื่องเคียงชุดเดียวกันก่อน (เพื่อรวมจำนวน)
        target = None
        for it in order.items.filter(menu=menu):
            existing = set(it.addons.values_list('pk', flat=True))
            if existing == addons_pk:
                target = it
                break

        if target:
            target.quantity += quantity
            if notes:
                target.notes = notes
            target.save()
            action = 'เพิ่มจำนวน'
        else:
            target = OrderItem.objects.create(order=order, menu=menu, quantity=quantity)
            if addons:
                target.addons.set(addons)
            if notes:
                target.notes = notes
            action = 'เพิ่ม'

        target.refresh_from_db()
        selected = str(menu.pk)
        addon_txt = ' (+เครื่องเคียง)' if addons else ''
        messages.success(request, f'{action} "{menu.name}" x{quantity}{addon_txt} ลงออเดอร์เรียบร้อย')
    else:
        for errors in form.errors.values():
            for error in errors:
                messages.error(request, error)

    url = reverse('order_detail', args=[order.pk])
    if back_page.isdigit() or selected:
        query = []
        if back_page.isdigit():
            query.append(f'page={back_page}')
        if selected:
            query.append(f'selected={selected}')
        url += '?' + '&'.join(query)
    url += '#menu-picker'
    return redirect(url)


@require_POST
@login_required
def order_item_qty(request, item_pk):
    """เพิ่ม/ลดจำนวนของรายการ (delta = 1 หรือ -1)"""
    item = get_object_or_404(OrderItem.objects.select_related('order', 'menu'), pk=item_pk)
    if not _user_can_access_order(request.user, item.order):
        messages.error(request, 'คุณไม่มีสิทธิ์แก้ไขออเดอร์นี้')
        return redirect('order_list')
    back_page = request.POST.get('page', '') or request.GET.get('page', '')
    try:
        delta = int(request.POST.get('delta', 0))
    except (TypeError, ValueError):
        delta = 0

    new_qty = item.quantity + delta
    if delta == 0:
        pass
    elif new_qty <= 0:
        name = item.menu.name
        item.delete()
        messages.warning(request, f'ลบ "{name}" ออกจากออเดอร์ (จำนวนเหลือ 0)')
    else:
        item.quantity = new_qty
        item.save()

    url = reverse('order_detail', args=[item.order.pk])
    if back_page.isdigit():
        url += f'?page={back_page}'
    url += '#cart-list'
    return redirect(url)


@require_POST
@login_required
def order_item_remove(request, item_pk):
    """ลบรายการเมนูออกจากออเดอร์"""
    item = get_object_or_404(OrderItem.objects.select_related('order'), pk=item_pk)
    if not _user_can_access_order(request.user, item.order):
        messages.error(request, 'คุณไม่มีสิทธิ์แก้ไขออเดอร์นี้')
        return redirect('order_list')
    back_page = request.POST.get('page', '') or request.GET.get('page', '')
    order_pk = item.order.pk
    name = item.menu.name
    item.delete()
    messages.warning(request, f'ลบ "{name}" ออกจากออเดอร์เรียบร้อย')
    url = reverse('order_detail', args=[order_pk])
    if back_page.isdigit():
        url += f'?page={back_page}'
    url += '#cart-list'
    return redirect(url)


@login_required
def order_delete(request, pk):
    """ลบทั้งออเดอร์ (มีหน้ายืนยันก่อนลบ)"""
    order = get_object_or_404(Order, pk=pk)
    if not _user_can_access_order(request.user, order):
        messages.error(request, 'คุณไม่มีสิทธิ์ลบออเดอร์นี้')
        return redirect('order_list')
    if request.method == 'POST':
        label = str(order)
        order.delete()
        messages.warning(request, f'ลบ {label} เรียบร้อยแล้ว')
        return redirect('order_list')
    return render(request, 'order/order_confirm_delete.html', {'order': order})


@staff_required
def order_pay(request, pk):
    """หน้าชำระเงิน: เลือกวิธีจ่าย (รวม QR PromptPay ถ้าร้านตั้งค่าไว้) แล้วปิดบิลเป็น 'ชำระแล้ว'"""
    order = get_object_or_404(Order, pk=pk)
    if not _user_can_access_order(request.user, order):
        messages.error(request, 'คุณไม่มีสิทธิ์เข้าถึงออเดอร์นี้')
        return redirect('order_list')

    if order.is_paid or order.status in ('refunded', 'completed'):
        messages.info(request, f'{order} ถูกปิดบิลแล้ว (สถานะ: {order.get_status_display()})')
        return redirect('order_detail', pk=order.pk)

    if not order.items.exists():
        messages.warning(request, 'ออเดอร์นี้ยังไม่มีรายการอาหาร ไม่สามารถชำระเงินได้')
        return redirect('order_detail', pk=order.pk)

    delivery = getattr(order, 'delivery', None)
    if order.order_type == 'delivery' and delivery is None:
        messages.warning(request, 'ออเดอร์เดลิเวอร์รี่ยังไม่ได้กรอกข้อมูลจัดส่ง รบกวนเพิ่มก่อนรับชำระเงิน')
        return redirect('order_detail', pk=order.pk)

    if request.method == 'POST':
        method = request.POST.get('method', '')
        if method not in dict(Order.PAYMENT_METHODS):
            messages.error(request, 'กรุณาเลือกวิธีชำระเงินก่อนยืนยัน')
        else:
            # การนับจำนวนครั้งการใช้คูปอง + สะสมคะแนนสมาชิก ระบบจัดการให้ใน Order signal
            # (เฉพาะตอนปิดบิลครั้งแรก กัน save ซ้ำแล้วนับซ้ำ)
            order.payment_method = method
            order.paid_at = timezone.now()
            order.status = 'paid'
            order.save()
            messages.success(
                request,
                f'รับชำระ {order.discounted_total} บาท '
                f'({order.get_payment_method_display}) เรียบร้อย ขอบคุณที่มาใช้บริการ 🙏'
            )
            return redirect('order_detail', pk=order.pk)

    return render(request, 'order/order_pay.html', {
        'order': order,
        'qr_data': _promptpay_qr_for(order),
        'qr_enabled': bool(StoreLocation.get_setting().promptpay_id),
    })


@login_required
def order_pay_qr(request, pk):
    """หน้าชำระเงินด้วย QR PromptPay สำหรับลูกค้า (สแกน + แจ้งชำระ + แนบหลักฐาน)
    การปิดบิลจริงยังต้องให้พนักงานยืนยันในหน้า order_pay เพื่อป้องกันการปลอมแปลง
    """
    order = get_object_or_404(Order, pk=pk)
    if not _user_can_access_order(request.user, order):
        messages.error(request, 'คุณไม่มีสิทธิ์เข้าถึงออเดอร์นี้')
        return redirect('order_list')

    if order.is_paid or order.status in ('cancelled', 'refunded'):
        messages.info(request, f'{order} ถูกปิดบิลแล้ว (สถานะ: {order.get_status_display()})')
        return redirect('order_detail', pk=order.pk)

    if not order.items.exists():
        messages.warning(request, 'ออเดอร์นี้ยังไม่มีรายการอาหาร ไม่สามารถชำระเงินได้')
        return redirect('order_detail', pk=order.pk)

    store = StoreLocation.get_setting()
    qr_enabled = bool(store.promptpay_id)
    if request.method == 'POST':
        method = (
            request.POST.get('method')
            or request.POST.get('payment_method')
            or request.POST.get('pay_method')
            or ''
        ).strip()
        # ลูกค้าสามารถเลือกวิธีชำระได้หลายทาง (ไม่จำกัดเฉพาะ QR)
        if method not in dict(Order.PAYMENT_METHODS):
            messages.error(request, 'กรุณาเลือกวิธีชำระเงินก่อนยืนยัน')
            return redirect('order_pay_qr', pk=order.pk)
        if method == 'qr' and not qr_enabled:
            messages.error(request, 'ร้านยังไม่ได้ตั้งค่า PromptPay ID จึงไม่สามารถชำระด้วย QR ได้')
            return redirect('order_pay_qr', pk=order.pk)
        form = CustomerQRPaymentForm(
            request.POST, request.FILES, instance=order,
            require_slip=(method == 'transfer'),
        )
        if form.is_valid():
            order = form.save(commit=False)
            order.payment_method = method
            order.qr_payment_requested_at = timezone.now()
            order.save()
            # แจ้งพนักงาน/ระบบ
            try:
                from .automation import notify
                notify(
                    'payment',
                    f'แจ้งชำระ {order.get_payment_method_display()} {order.display_code}',
                    f'{order.table_display} ยอด {order.discounted_total} บาท ลูกค้าแจ้งชำระผ่าน '
                    f'{order.get_payment_method_display()} แล้ว'
                    + (f' (อ้างอิง: {order.qr_payment_reference})' if order.qr_payment_reference else ''),
                    f'/orders/{order.pk}/pay/',
                )
            except Exception:
                pass
            messages.success(
                request,
                'บันทึกการแจ้งชำระแล้ว พนักงานจะตรวจสอบและปิดบิลให้ในเร็ว ๆ นี้ 🙏'
            )
            return redirect('order_detail', pk=order.pk)

    if request.method != 'POST':
        form = CustomerQRPaymentForm(instance=order)

    return render(request, 'order/order_pay_qr.html', {
        'order': order,
        'store': store,
        'qr_data': _promptpay_qr_for(order) if qr_enabled else None,
        'qr_enabled': qr_enabled,
        'form': form,
    })


@staff_required
def order_receipt(request, pk):
    """หน้าใบเสร็จรับเงิน (สามารถพิมพ์ผ่านปุ่ม Print)"""
    order = get_object_or_404(
        Order.objects.prefetch_related('items', 'items__menu', 'items__addons')
        .select_related('created_by', 'delivery'),
        pk=pk,
    )
    if not _user_can_access_order(request.user, order):
        messages.error(request, 'คุณไม่มีสิทธิ์เข้าถึงออเดอร์นี้')
        return redirect('order_list')
    store = StoreLocation.get_setting()
    return render(request, 'order/order_receipt.html', {'order': order, 'store': store})


@login_required
def table_qr(request, table_number):
    """หน้า QR code สำหรับสั่งอาหารประจำโต๊ะ"""
    order_url = request.build_absolute_uri(
        reverse('order_list') + f'?table={table_number}'
    )
    return render(request, 'order/table_qr.html', {
        'table_number': table_number,
        'order_url': order_url,
    })


@staff_required
def export_orders_csv(request):
    """ส่งออกรายการขาย (ยอดชำระแล้ว) เป็นไฟล์ CSV"""
    response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    response['Content-Disposition'] = (
        'attachment; filename="sales_report.csv"'
    )
    writer = csv.writer(response)
    writer.writerow([
        'ออเดอร์', 'ประเภท', 'โต๊ะ/ผู้รับ', 'เบอร์โทร', 'ที่อยู่', 'ค่าส่ง', 'สถานะส่ง',
        'เวลาชำระ', 'วิธีชำระ', 'เมนู', 'จำนวน', 'ราคา/หน่วย', 'รวม', 'สั่งโดย',
    ])
    for row in _sale_rows():
        writer.writerow([
            row['order_id'], row['order_type'], row['recipient'] if row['order_type'] == 'เดลิเวอร์รี่' else row['table_number'],
            row['phone'], row['address'], row['delivery_fee'], row['delivery_status'],
            row['paid_at'], row['payment_method'], row['menu_name'], row['quantity'],
            row['unit_price'], row['subtotal'], row['created_by'],
        ])
    return response


@staff_required
def export_orders_xlsx(request):
    """ส่งออกรายการขาย (ยอดชำระแล้ว) เป็นไฟล์ Excel (.xlsx)"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'รายงานยอดขาย'

    headers = [
        'ออเดอร์', 'ประเภท', 'โต๊ะ/ผู้รับ', 'เบอร์โทร', 'ที่อยู่', 'ค่าส่ง', 'สถานะส่ง',
        'เวลาชำระ', 'วิธีชำระ', 'เมนู', 'จำนวน', 'ราคา/หน่วย', 'รวม', 'สั่งโดย',
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='FF5E62')
    ws.freeze_panes = 'A2'

    for row in _sale_rows():
        ws.append([
            row['order_id'], row['order_type'],
            row['recipient'] if row['order_type'] == 'เดลิเวอร์รี่' else row['table_number'],
            row['phone'], row['address'], row['delivery_fee'], row['delivery_status'],
            row['paid_at'], row['payment_method'], row['menu_name'], row['quantity'],
            row['unit_price'], row['subtotal'], row['created_by'],
        ])

    # สรุปยอดรวมท้ายตาราง
    last = ws.max_row
    ws.append([
        'รวมทั้งหมด', '', '', '', '', '', '',
        '', '', '', '', '', f"=SUM(M{2}:M{last})", '',
    ])
    ws.cell(row=last + 1, column=13).font = Font(bold=True)

    for col_idx in range(1, len(headers) + 1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = 16

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="sales_report.xlsx"'
    wb.save(response)
    return response


@require_POST
@staff_required
def order_update_status(request, pk):
    """เปลี่ยนสถานะออเดอร์ (รองรับออเดอร์เดลิเวอร์รี่ด้วย)"""
    order = get_object_or_404(Order, pk=pk)
    if not _user_can_access_order(request.user, order):
        messages.error(request, 'คุณไม่มีสิทธิ์แก้ไขออเดอร์นี้')
        return redirect('order_list')
    new_status = request.POST.get('status', '')
    valid_transitions = {
        'open': ['preparing', 'cancelled', 'pending'],
        'preparing': ['ready', 'cancelled'],
        'ready': ['served', 'cancelled'],
        'served': ['paid', 'completed'],
        'paid': ['refunded'],
        'cancelled': [],
        'pending': ['confirmed', 'cancelled'],
        'confirmed': ['preparing', 'cancelled'],
        'completed': ['refunded'],
        'refunded': [],
    }
    allowed = valid_transitions.get(order.status, [])
    if new_status in allowed:
        order.status = new_status
        if new_status == 'paid':
            order.paid_at = timezone.now()
        order.save()

        # ออเดอร์เดลิเวอร์รี่: สถานะ 'served' = จัดส่งสำเร็จ → อัปเดตสถานะการจัดส่งพร้อมกัน
        if order.order_type == 'delivery' and new_status in ('served', 'ready', 'preparing'):
            delivery = getattr(order, 'delivery', None)
            if delivery is not None:
                if new_status == 'served':
                    delivery.status = 'delivered'
                    delivery.delivered_at = timezone.now()
                elif new_status in ('ready', 'preparing'):
                    delivery.status = 'pending'
                    delivery.delivered_at = None
                delivery.save()

        status_display = order.get_status_display()
        messages.success(request, f'เปลี่ยนสถานะออเดอร์ {order.display_code} เป็น "{status_display}" เรียบร้อย')
    else:
        messages.error(request, 'ไม่สามารถเปลี่ยนสถานะได้')
    return redirect('order_detail', pk=order.pk)
